"""Generate the founder bulk-post Excel template (Scorred_Posts_Template.xlsx).

The Reference sheet (valid handles, communities, catalogue SKUs) is pulled LIVE
from the target database, so regenerate the file whenever those change:

    backend/.venv/bin/python backend/scripts/make_bulk_post_template.py
    # target defaults to QA via backend/.prod.env; override with DATABASE_URL

Output: <repo>/bulk_upload/Scorred_Posts_Template.xlsx
Founders fill the Posts sheet (one row = one post) and send the file back;
ingest with backend/scripts/ingest_bulk_posts.py.
"""
import asyncio
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bulk_common import connect, resolve_db_url  # noqa: E402

OUT = Path(__file__).resolve().parents[2] / "bulk_upload" / "Scorred_Posts_Template.xlsx"

TYPES = ["showcase", "discussion", "review", "poll", "iso"]
CATEGORIES = ["figures", "designer", "kits", "diecast", "tcg"]
ISO_CONDITIONS = ["any", "sealed", "mib", "bib", "loose"]

HEADERS = [
    ("author_handle", 18, "Who posts it. Pick from the dropdown (real Scorred accounts)."),
    ("type", 12, "showcase / discussion / review / poll / iso"),
    ("title", 28, "Optional headline shown above the post"),
    ("body", 60, "The post text. Required for every post."),
    ("image_url_1", 40, "Direct image link (or Google Drive share link)"),
    ("image_url_2", 40, ""),
    ("image_url_3", 40, ""),
    ("image_url_4", 40, ""),
    ("category", 12, "Optional. Which collector category the post belongs to"),
    ("tags", 24, "Optional. Comma-separated, e.g. HotToys, Grails"),
    ("community", 22, "Optional. Post into this community (dropdown)"),
    ("review_sku", 22, "REVIEW ONLY (required): the catalogue entry being reviewed — see Reference sheet"),
    ("review_rating", 13, "REVIEW ONLY (required): 1-5"),
    ("poll_option_1", 20, "POLL ONLY: at least options 1 and 2"),
    ("poll_option_2", 20, ""),
    ("poll_option_3", 20, ""),
    ("poll_option_4", 20, ""),
    ("iso_item", 30, "ISO ONLY (required): what you're looking for"),
    ("iso_budget_inr", 14, "ISO optional: max budget in whole rupees"),
    ("iso_condition", 13, "ISO optional: any/sealed/mib/bib/loose"),
    ("iso_city", 16, "ISO optional: city, or blank = Anywhere in India"),
    ("days_ago", 10, "Optional 0-30: backdate the post to look organic"),
]

INSTRUCTIONS = [
    ("Scorred — Bulk Post Upload Template", "title"),
    (f"Generated {date.today().isoformat()}. One row on the Posts sheet = one post on Scorred.", ""),
    ("", ""),
    ("HOW TO USE", "h"),
    ("1. Fill one row per post on the 'Posts' sheet. Dropdowns show only valid values.", ""),
    ("2. body is always required. author_handle and type are always required.", ""),
    ("3. Send the file back to RC — posts are checked first, and you get a report of "
     "anything that needs fixing before anything goes live.", ""),
    ("", ""),
    ("POST TYPES", "h"),
    ("showcase — 'look at this' post. Needs at least one image. Body describes the piece.", ""),
    ("discussion — a text post / question to the community. Images optional.", ""),
    ("review — your verdict on a specific collectible. REQUIRES review_sku (pick the exact "
     "item from the Reference sheet) and review_rating 1-5. Images strongly recommended.", ""),
    ("poll — a question with 2-4 answer options (poll_option_1..4). No images needed.", ""),
    ("iso — 'In Search Of' / wanted post. REQUIRES iso_item. Budget, condition, city optional.", ""),
    ("", ""),
    ("IMAGES", "h"),
    ("Paste direct image links (ending .jpg/.png/.webp) OR a Google Drive share link "
     "(set sharing to 'Anyone with the link'). Up to 4 per post.", ""),
    ("Images are downloaded and re-hosted on Scorred's own storage during ingest — "
     "the link only needs to work once.", ""),
    ("Use photos you took yourself, or images you have the right to use.", ""),
    ("", ""),
    ("REFERENCE SHEET", "h"),
    ("Valid account handles, communities, and the catalogue (for review_sku) are listed on "
     "the Reference sheet. Don't edit that sheet — regenerate the template to refresh it.", ""),
]


async def fetch_reference():
    conn = await connect(resolve_db_url())
    handles = await conn.fetch(
        "SELECT handle, name FROM users WHERE NOT is_suspended ORDER BY handle"
    )
    communities = await conn.fetch(
        "SELECT id, name FROM communities ORDER BY name"
    )
    catalogue = await conn.fetch(
        "SELECT sku, title, brand, category FROM catalogue "
        "WHERE status != 'removed' ORDER BY category, brand, title"
    )
    await conn.close()
    return handles, communities, catalogue


def build(handles, communities, catalogue):
    wb = Workbook()

    head_fill = PatternFill("solid", fgColor="FF2442")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    hint_font = Font(color="777777", italic=True, size=9)
    thin = Border(bottom=Side(style="thin", color="DDDDDD"))

    # ── Instructions ─────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 110
    for i, (text, style) in enumerate(INSTRUCTIONS, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if style == "title":
            c.font = Font(bold=True, size=16, color="FF2442")
        elif style == "h":
            c.font = Font(bold=True, size=12)

    # ── Posts (entry sheet) ──────────────────────────────────────────────
    ws = wb.create_sheet("Posts")
    for col, (name, width, hint) in enumerate(HEADERS, start=1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = width
        h = ws.cell(row=1, column=col, value=name)
        h.fill, h.font = head_fill, head_font
        h.alignment = Alignment(horizontal="center")
        hc = ws.cell(row=2, column=col, value=hint)
        hc.font = hint_font
        hc.alignment = Alignment(wrap_text=True, vertical="top")
        hc.border = thin
    ws.freeze_panes = "C3"
    ws.row_dimensions[2].height = 42

    n_rows = 202  # generous entry range for validations
    col_of = {name: i + 1 for i, (name, _, _) in enumerate(HEADERS)}

    def add_dv(formula, col_name, *, allow_blank=True, err="Pick a value from the dropdown."):
        dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank,
                            showErrorMessage=True, error=err)
        letter = get_column_letter(col_of[col_name])
        dv.add(f"{letter}3:{letter}{n_rows}")
        ws.add_data_validation(dv)

    add_dv(f'"{",".join(TYPES)}"', "type", err="Must be one of: " + ", ".join(TYPES))
    add_dv(f'"{",".join(CATEGORIES)}"', "category")
    add_dv(f'"{",".join(ISO_CONDITIONS)}"', "iso_condition")
    add_dv('"1,2,3,4,5"', "review_rating", err="Rating is 1-5")
    if handles:
        add_dv(f"Reference!$A$3:$A${2 + len(handles)}", "author_handle",
               err="Must be an existing Scorred handle — see Reference sheet")
    if communities:
        add_dv(f"Reference!$D$3:$D${2 + len(communities)}", "community")
    if catalogue:
        add_dv(f"Reference!$G$3:$G${2 + len(catalogue)}", "review_sku",
               err="Must be a catalogue SKU — see Reference sheet")
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="30",
                        allow_blank=True, showErrorMessage=True, error="0-30 days")
    dv.add(f"{get_column_letter(col_of['days_ago'])}3:{get_column_letter(col_of['days_ago'])}{n_rows}")
    ws.add_data_validation(dv)

    # ── Reference ────────────────────────────────────────────────────────
    ws = wb.create_sheet("Reference")
    blocks = [
        ("A", "ACCOUNTS", ["handle", "name"],
         [(r["handle"], r["name"]) for r in handles]),
        ("D", "COMMUNITIES", ["id", "name"],
         [(r["id"], r["name"]) for r in communities]),
        ("G", "CATALOGUE (for review_sku)", ["sku", "title", "brand", "category"],
         [(r["sku"], r["title"], r["brand"], r["category"]) for r in catalogue]),
    ]
    for start_letter, title, cols, rows in blocks:
        start = ws[f"{start_letter}1"].column
        t = ws.cell(row=1, column=start, value=title)
        t.font = Font(bold=True, size=12, color="FF2442")
        for j, name in enumerate(cols):
            c = ws.cell(row=2, column=start + j, value=name)
            c.fill, c.font = head_fill, head_font
        for i, row in enumerate(rows, start=3):
            for j, val in enumerate(row):
                ws.cell(row=i, column=start + j, value=val)
    for letter, width in (("A", 22), ("B", 24), ("D", 24), ("E", 30),
                          ("G", 18), ("H", 46), ("I", 18), ("J", 12)):
        ws.column_dimensions[letter].width = width

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"✓ wrote {OUT}")
    print(f"  reference: {len(handles)} accounts, {len(communities)} communities, "
          f"{len(catalogue)} catalogue entries")


if __name__ == "__main__":
    h, c, cat = asyncio.run(fetch_reference())
    build(h, c, cat)
