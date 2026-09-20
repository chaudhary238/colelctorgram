"""Ingest a filled Scorred_Posts_Template.xlsx through the LIVE API.

    # 1. Validate only (default — nothing is written):
    backend/.venv/bin/python backend/scripts/ingest_bulk_posts.py <file.xlsx>

    # 2. Apply (logs in as each author and posts via the API):
    backend/.venv/bin/python backend/scripts/ingest_bulk_posts.py <file.xlsx> --apply

Why the API and not direct DB inserts: posting through the API awards XP,
fires notifications and keeps every counter honest — the feed, leaderboard
and rewards behave exactly as if the author had used the app (2026-09-20
founder decision, replacing the direct-DB ingest).

Author credentials: seed personas (@seed.scorred.com) use `seed_password`
from backend/.prod.env; the collectorhub admin uses `seed_admin_password`.
Rows authored by anyone else are skipped with a clear message — real people
post via the app.

Other behaviour:
- API base = qa_api_base in .prod.env (default https://api.qa.scorred.com/v1).
- Every image URL is verified to actually serve an image. With r2_* keys in
  .prod.env images are re-hosted on R2 first; without them the source URL is
  stored directly (fine for hotlink-friendly hosts; Google Drive links
  REQUIRE the R2 keys — Drive blocks hotlinking).
- days_ago: the post is created via the API (stamped 'now'), then created_at
  is backdated with one cosmetic DB update. XP keeps its awarded value.
- An audit manifest is written next to the input file.
"""
import asyncio
import json
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import httpx
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bulk_common import (  # noqa: E402
    connect, download_image, load_prod_env, normalize_image_url, r2_configured,
    resolve_db_url, upload_to_r2,
)

TYPES = {"showcase", "discussion", "review", "poll", "iso"}
CATEGORIES = {"figures", "designer", "kits", "diecast", "tcg"}
ISO_CONDITIONS = {"any", "sealed", "mib", "bib", "loose"}

COLS = [
    "author_handle", "type", "title", "body",
    "image_url_1", "image_url_2", "image_url_3", "image_url_4",
    "category", "tags", "community", "review_sku", "review_rating",
    "poll_option_1", "poll_option_2", "poll_option_3", "poll_option_4",
    "iso_item", "iso_budget_inr", "iso_condition", "iso_city", "days_ago",
]


class ApiClient:
    """Login-per-author API poster for the seed/admin accounts."""

    def __init__(self):
        env = load_prod_env()
        self.base = env.get("qa_api_base", "https://api.qa.scorred.com/v1").rstrip("/")
        self.persona_pw = env.get("seed_password")
        self.admin_pw = env.get("seed_admin_password")
        self.client = httpx.Client(base_url=self.base, timeout=90)
        self.tokens: dict[str, str] = {}
        self.seed_handles: set[str] = set()

    def password_for(self, handle: str) -> str | None:
        if handle == "collectorhub":
            return self.admin_pw
        if handle in self.seed_handles:
            return self.persona_pw
        return None

    def login(self, handle: str) -> str | None:
        if handle in self.tokens:
            return self.tokens[handle]
        pw = self.password_for(handle)
        if not pw:
            return None
        r = self.client.post("/auth/login", json={"identifier": handle, "password": pw})
        if r.status_code != 200:
            return None
        self.tokens[handle] = r.json()["access_token"]
        return self.tokens[handle]

    def create_post(self, handle: str, payload: dict) -> dict:
        token = self.login(handle)
        if not token:
            raise RuntimeError(
                f"no credential for @{handle} (only seed personas + collectorhub "
                "can be posted for — real accounts post via the app)")
        for attempt in range(4):
            r = self.client.post("/posts", json=payload,
                                 headers={"Authorization": f"Bearer {token}"})
            if r.status_code == 429:
                time.sleep(5 * (attempt + 1))
                continue
            break
        if r.status_code not in (200, 201):
            raise RuntimeError(f"POST /posts -> {r.status_code}: {r.text[:200]}")
        time.sleep(0.12)
        return r.json()


def read_rows(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    if "Posts" not in wb.sheetnames:
        raise SystemExit("No 'Posts' sheet in workbook")
    ws = wb["Posts"]
    header = [c.value for c in ws[1]]
    if header[: len(COLS)] != COLS:
        raise SystemExit(
            "Posts sheet columns don't match the template — regenerate the "
            "template and copy your rows over.\n"
            f"expected: {COLS}\nfound:    {header[:len(COLS)]}"
        )
    rows = []
    for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        vals = {k: (str(v).strip() if v is not None and str(v).strip() != "" else None)
                for k, v in zip(COLS, row)}
        if any(vals.values()):
            vals["_row"] = i
            rows.append(vals)
    return rows


def validate(r: dict, handles: set, communities: set, skus: set) -> list[str]:
    errs = []
    if not r["author_handle"]:
        errs.append("author_handle missing")
    elif r["author_handle"] not in handles:
        errs.append(f"unknown handle '{r['author_handle']}'")
    t = (r["type"] or "").lower()
    if t not in TYPES:
        errs.append(f"type must be one of {sorted(TYPES)}")
    if not r["body"]:
        errs.append("body missing")
    imgs = [r[f"image_url_{i}"] for i in range(1, 5) if r[f"image_url_{i}"]]
    if t == "showcase" and not imgs:
        errs.append("showcase needs at least one image")
    if r["category"] and r["category"].lower() not in CATEGORIES:
        errs.append(f"category must be one of {sorted(CATEGORIES)}")
    if r["community"] and r["community"] not in communities:
        errs.append(f"unknown community '{r['community']}'")
    if t == "review":
        if not r["review_sku"]:
            errs.append("review needs review_sku")
        elif r["review_sku"] not in skus:
            errs.append(f"review_sku '{r['review_sku']}' not in catalogue")
        try:
            if int(float(r["review_rating"] or 0)) not in range(1, 6):
                raise ValueError
        except (TypeError, ValueError):
            errs.append("review_rating must be 1-5")
    if t == "poll":
        opts = [r[f"poll_option_{i}"] for i in range(1, 5) if r[f"poll_option_{i}"]]
        if len(opts) < 2:
            errs.append("poll needs at least 2 options")
        if len(set(opts)) != len(opts):
            errs.append("poll options must be unique")
    if t == "iso" and not r["iso_item"]:
        errs.append("iso needs iso_item")
    if r["iso_condition"] and r["iso_condition"].lower() not in ISO_CONDITIONS:
        errs.append(f"iso_condition must be one of {sorted(ISO_CONDITIONS)}")
    if r["iso_budget_inr"]:
        try:
            if float(r["iso_budget_inr"]) <= 0:
                raise ValueError
        except (TypeError, ValueError):
            errs.append("iso_budget_inr must be a positive number of rupees")
    if r["days_ago"]:
        try:
            if not 0 <= int(float(r["days_ago"])) <= 30:
                raise ValueError
        except (TypeError, ValueError):
            errs.append("days_ago must be 0-30")
    return errs


def norm_tags(raw: str | None) -> list[str]:
    if not raw:
        return []
    out = []
    for part in raw.split(","):
        part = part.strip().lstrip("#").replace(" ", "")
        if part:
            out.append(f"#{part}")
    return out


def build_payload(r: dict, images: list[str]) -> dict:
    t = r["type"].lower()
    payload = dict(
        type=t, body=r["body"], title=r["title"], images=images,
        tags=norm_tags(r["tags"]),
        category=r["category"].lower() if r["category"] else None,
        community_id=r["community"], to_feed=True,
    )
    if t == "review":
        payload |= dict(ref_sku=r["review_sku"],
                        review_rating=int(float(r["review_rating"])))
    if t == "poll":
        payload |= dict(poll_options={
            r[f"poll_option_{i}"]: 0 for i in range(1, 5) if r[f"poll_option_{i}"]})
    if t == "iso":
        cond = (r["iso_condition"] or "").lower()
        payload |= dict(
            iso_item=r["iso_item"],
            iso_budget=int(float(r["iso_budget_inr"]) * 100) if r["iso_budget_inr"] else None,
            iso_condition=cond if cond and cond != "any" else None,
            iso_city=r["iso_city"])
    return payload


def resolve_images(r: dict, rehost: bool, manifest: dict) -> list[str]:
    images = []
    for i in range(1, 5):
        src = r[f"image_url_{i}"]
        if not src:
            continue
        data, ctype = download_image(src)  # validates it IS an image
        hosted = upload_to_r2(data, ctype, prefix="posts") if rehost \
            else normalize_image_url(src)
        images.append(hosted)
        manifest["images"].append({"source": src, "hosted": hosted})
    return images


async def main():
    args = list(sys.argv[1:])
    apply = "--apply" in args
    files = [a for a in args if not a.startswith("--")]
    if len(files) != 1:
        raise SystemExit(__doc__)
    path = Path(files[0]).expanduser()
    rows = read_rows(path)
    print(f"{len(rows)} filled row(s) in {path.name}")

    conn = await connect(resolve_db_url())
    seed_handles = {r["handle"] for r in await conn.fetch(
        "SELECT handle FROM users WHERE email LIKE '%@seed.scorred.com' AND NOT is_suspended")}
    all_handles = {r["handle"] for r in await conn.fetch(
        "SELECT handle FROM users WHERE NOT is_suspended")}
    communities = {r["id"] for r in await conn.fetch("SELECT id FROM communities")}
    skus = {r["sku"] for r in await conn.fetch(
        "SELECT sku FROM catalogue WHERE status != 'removed'")}

    valid, invalid = [], []
    for r in rows:
        errs = validate(r, all_handles, communities, skus)
        (invalid if errs else valid).append((r, errs))
    for r, errs in invalid:
        print(f"  ✗ row {r['_row']}: " + "; ".join(errs))
    print(f"valid={len(valid)} invalid={len(invalid)}")

    if not apply:
        postable = {h for h in {r["author_handle"] for r, _ in valid}
                    if h in seed_handles or h == "collectorhub"}
        skipped = {r["author_handle"] for r, _ in valid} - postable
        if skipped:
            print(f"note: rows by {sorted(skipped)} will be SKIPPED on apply "
                  "(no seed credential — those people post via the app)")
        print("\nDry run only — rerun with --apply to ingest the valid rows.")
        await conn.close()
        return

    api = ApiClient()
    api.seed_handles = seed_handles
    rehost = r2_configured()
    if not rehost:
        print("(no R2 keys in .prod.env — storing verified source URLs directly)")
    manifest = {"file": str(path), "ingested_at": datetime.now(timezone.utc).isoformat(),
                "rehosted": rehost, "via": "api", "images": [], "posts": []}
    now = datetime.now(timezone.utc)
    ok = 0
    backdates: list[tuple[str, datetime]] = []
    for r, _ in valid:
        try:
            images = resolve_images(r, rehost, manifest)
            if r["type"].lower() == "showcase" and not images:
                raise RuntimeError("no usable images")
            post = api.create_post(r["author_handle"], build_payload(r, images))
        except Exception as e:
            print(f"  ✗ row {r['_row']} skipped — {e}")
            continue
        days = int(float(r["days_ago"])) if r["days_ago"] else 0
        if days:
            backdates.append((post["id"], now - timedelta(days=days,
                                                          minutes=(r["_row"] * 7) % 600)))
        manifest["posts"].append({"row": r["_row"], "post_id": post["id"],
                                  "author": r["author_handle"], "type": r["type"].lower()})
        ok += 1
        print(f"  ✓ row {r['_row']}: {r['type'].lower()} by @{r['author_handle']} "
              f"({len(images)} img)")

    for pid, ts in backdates:
        await conn.execute(
            "UPDATE posts SET created_at=$1, updated_at=$1 WHERE id=$2::uuid", ts, pid)
    if backdates:
        print(f"backdated {len(backdates)} post(s)")
    await conn.close()

    mpath = path.with_suffix(path.suffix + ".ingest_manifest.json")
    mpath.write_text(json.dumps(manifest, indent=2))
    print(f"\n✓ ingested {ok}/{len(valid)} valid rows via API; manifest: {mpath.name}")


if __name__ == "__main__":
    asyncio.run(main())
